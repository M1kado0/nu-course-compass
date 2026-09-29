import hashlib
import io
import json
import sys
from pathlib import Path

import httpx
import pytest
from langchain_core.documents import Document
from docx import Document as WordDocument
from nu_course_compass.ingestion import import_moe_syllabi

from nu_course_compass.ingestion.import_moe_syllabi import (
    GOOGLE_DOC,
    GOOGLE_FOLDER,
    PDF,
    GoogleReadClient,
    drive_file_id,
    import_in_batches,
    import_rows,
)
from nu_course_compass.ingestion.load_syllabus_documents import load_syllabus_documents
from nu_course_compass.ingestion.moe_sheet import SyllabusIndexRow, parse_sheet_grid


def cell(value: str = "", link: str | None = None) -> dict:
    result = {"formattedValue": value}
    if link:
        result["hyperlink"] = link
    return result


def test_parse_repeated_school_headers_and_hyperlinks() -> None:
    grid = {"rowData": [
        {"values": [cell(), cell("Fall'22 Offline Syllabi Collection")]},
        {"values": [cell(), cell("SMG")]},
        {"values": [cell(), cell("Course Code"), cell("Instructor Name"),
                    cell("Semester"), cell("Link to syllabi")]},
        {"values": [cell(), cell("GEOL 201"), cell("Laurent Richard"),
                    cell("Fall'22"), cell("Syllabus", "https://drive.google.com/file/d/abc/view")]},
        {"values": [cell(), cell("SEDS")]},
        {"values": [cell(), cell("Course Code"), cell("Instructor Name"),
                    cell("Semester"), cell("Link to syllabi")]},
        {"values": [cell(), cell("CEE 200"), cell("Other"), cell("Fall 2022"),
                    cell("https://docs.google.com/document/d/xyz/edit")]},
        {"values": [cell(), cell("CEE 201"), cell("Other"), cell("Fall 2022"), cell()]},
    ]}
    rows, issues = parse_sheet_grid(sheet_title="Fall'22", sheet_gid=7, grid=grid)
    assert [(row.course_code, row.school, row.sheet_row) for row in rows] == [
        ("GEOL 201", "SMG", 4), ("CEE 200", "SEDS", 7)
    ]
    assert rows[0].syllabus_url == "https://drive.google.com/file/d/abc/view"
    assert rows[0].source_url.endswith("range=B4:E4")
    assert len(issues) == 1
    assert issues[0].sheet_row == 8


@pytest.mark.parametrize(("url", "file_id", "reason"), [
    ("https://drive.google.com/file/d/abc/view", "abc", ""),
    ("https://docs.google.com/document/d/xyz/edit", "xyz", ""),
    ("https://drive.google.com/drive/u/1/folders/f1", None, "folder_link"),
    ("https://evil.example/file/d/abc/view", None, "unsupported_url"),
])
def test_drive_file_id(url: str, file_id: str | None, reason: str) -> None:
    assert drive_file_id(url) == (file_id, reason)


def row(url: str, number: int = 4) -> SyllabusIndexRow:
    return SyllabusIndexRow("CSCI 361", "Instructor", "Fall'25", "SCAI", url,
                            "Fall'25", 1, number)


def docx_bytes() -> bytes:
    document = WordDocument()
    document.add_heading("Course overview", level=1)
    document.add_paragraph("Topics include computer vision.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Assessment"
    table.cell(0, 1).text = "Weight"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_google_client_uses_each_apis_documented_host() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    seen: list[str] = []

    def respond(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        return httpx.Response(200, json={})

    with GoogleReadClient(Credentials()) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(respond))
        client.get("/v4/spreadsheets/example")
        client.get("/drive/v3/files/example")

    assert seen == [
        "https://sheets.googleapis.com/v4/spreadsheets/example",
        "https://www.googleapis.com/drive/v3/files/example",
    ]


def test_google_client_retries_rate_limit_then_succeeds() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    attempts = 0
    waits: list[float] = []

    def respond(_: httpx.Request) -> httpx.Response:
        nonlocal attempts
        attempts += 1
        if attempts < 3:
            return httpx.Response(429)
        return httpx.Response(200, json={"ok": True})

    with GoogleReadClient(Credentials(), sleep=waits.append) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(respond))
        assert client.get("/drive/v3/files/example").json() == {"ok": True}

    assert attempts == 3
    assert waits == [1.0, 2.0]


def test_google_client_honors_retry_after_and_stops_after_four_attempts() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    attempts = 0
    waits: list[float] = []

    def respond(_: httpx.Request) -> httpx.Response:
        nonlocal attempts
        attempts += 1
        return httpx.Response(429, headers={"Retry-After": "3"})

    with GoogleReadClient(Credentials(), sleep=waits.append) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(respond))
        with pytest.raises(httpx.HTTPStatusError) as error:
            client.get("/drive/v3/files/example")

    assert error.value.response.status_code == 429
    assert attempts == 4
    assert waits == [3.0, 3.0, 3.0]


def test_google_client_does_not_retry_access_denied() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    attempts = 0

    def respond(_: httpx.Request) -> httpx.Response:
        nonlocal attempts
        attempts += 1
        return httpx.Response(403)

    with GoogleReadClient(Credentials(), sleep=lambda _: None) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(respond))
        with pytest.raises(httpx.HTTPStatusError):
            client.get("/drive/v3/files/example")

    assert attempts == 1


def test_google_client_downloads_and_validates_docx_blob() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    content = docx_bytes()
    requested: list[str] = []

    def respond(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        return httpx.Response(200, content=content)

    with GoogleReadClient(Credentials()) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(respond))
        assert client.download_docx("word") == content

    assert requested == ["https://www.googleapis.com/drive/v3/files/word?alt=media"]


def test_google_client_rejects_non_docx_download() -> None:
    class Credentials:
        valid = True
        token = "test-token"

    with GoogleReadClient(Credentials()) as client:
        client.http.close()
        client.http = httpx.Client(transport=httpx.MockTransport(
            lambda _: httpx.Response(200, content=b"<html>Access denied</html>")))
        with pytest.raises(ValueError, match="non-DOCX"):
            client.download_docx("word")


class FakeGoogleClient:
    def __init__(self, mime: str = PDF, downloadable: bool = True) -> None:
        self.mime = mime
        self.downloadable = downloadable
        self.download_calls = 0

    def file_metadata(self, file_id: str) -> dict:
        return {"id": file_id, "mimeType": self.mime,
                "modifiedTime": "2026-01-01T00:00:00Z",
                "capabilities": {"canDownload": self.downloadable}}

    def download_pdf(self, file_id: str, mime_type: str) -> bytes:
        self.download_calls += 1
        return b"%PDF-1.7\nfixture"


def test_import_is_idempotent_and_preserves_raw_bytes(tmp_path: Path) -> None:
    client = FakeGoogleClient()
    source = row("https://drive.google.com/file/d/abc/view")
    first = import_rows([source], client, tmp_path)
    second = import_rows([source], client, tmp_path)
    assert first[0]["status"] == "downloaded"
    assert second[0]["sha256"] == first[0]["sha256"]
    assert client.download_calls == 1
    assert len((tmp_path / "manifest.jsonl").read_text().splitlines()) == 1
    assert (tmp_path / first[0]["local_path"]).read_bytes() == b"%PDF-1.7\nfixture"


def test_shared_syllabus_downloads_once_but_preserves_both_rows(tmp_path: Path) -> None:
    client = FakeGoogleClient()
    shared_url = "https://drive.google.com/file/d/shared/view"

    results = import_in_batches([row(shared_url, 4), row(shared_url, 5)],
                                client, tmp_path, batch_size=1, batch_delay=0,
                                sleep=lambda _: None)

    assert client.download_calls == 1
    assert [result["sheet_row"] for result in results] == [4, 5]
    assert results[1]["reused_file"] is True
    assert len((tmp_path / "manifest.jsonl").read_text().splitlines()) == 2


def test_import_preserves_original_docx_snapshot(tmp_path: Path) -> None:
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    content = docx_bytes()

    class DocxClient(FakeGoogleClient):
        def download_docx(self, file_id: str) -> bytes:
            self.download_calls += 1
            return content

    result = import_rows([row("https://drive.google.com/file/d/word/view")],
                         DocxClient(mime), tmp_path)[0]

    assert result["status"] == "downloaded"
    assert result["content_type"] == mime
    assert str(result["local_path"]).endswith(".docx")
    assert (tmp_path / str(result["local_path"])).read_bytes() == content


def test_import_in_batches_checkpoints_and_pauses_between_batches(tmp_path: Path) -> None:
    client = FakeGoogleClient()
    rows = [row(f"https://drive.google.com/file/d/file{number}/view", number)
            for number in range(4, 9)]
    waits: list[float] = []

    results = import_in_batches(rows, client, tmp_path, batch_size=2,
                                batch_delay=1.5, sleep=waits.append)

    assert len(results) == 5
    assert client.download_calls == 5
    assert waits == [1.5, 1.5]
    assert len((tmp_path / "manifest.jsonl").read_text().splitlines()) == 5


def test_cli_all_rows_ignores_pilot_limit_and_uses_batches(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    rows = [row(f"https://drive.google.com/file/d/file{number}/view", number)
            for number in range(4, 15)]

    class Client(FakeGoogleClient):
        def __enter__(self) -> "Client":
            return self

        def __exit__(self, *_: object) -> None:
            pass

        def sheet_rows(self) -> tuple[list[SyllabusIndexRow], list[dict]]:
            return rows, []

    monkeypatch.setattr(import_moe_syllabi, "google_credentials", lambda *_: object())
    monkeypatch.setattr(import_moe_syllabi, "GoogleReadClient", lambda _: Client())
    monkeypatch.setattr(sys, "argv", ["import_moe_syllabi", "--client-secret", "unused.json",
                                   "--output", str(tmp_path), "--all-rows",
                                   "--batch-size", "5", "--batch-delay", "0"])

    import_moe_syllabi.main()

    assert json.loads(capsys.readouterr().out)["downloaded"] == 11
    assert len((tmp_path / "manifest.jsonl").read_text().splitlines()) == 11


def test_cli_rejects_all_rows_with_limit_before_auth(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(sys, "argv", ["import_moe_syllabi", "--client-secret", "unused.json",
                                   "--all-rows", "--limit", "20"])
    monkeypatch.setattr(import_moe_syllabi, "google_credentials",
                        lambda *_: pytest.fail("must reject before OAuth"))

    with pytest.raises(SystemExit) as error:
        import_moe_syllabi.main()

    assert error.value.code == 2


def test_folder_and_download_disabled_are_reported(tmp_path: Path) -> None:
    folder = FakeGoogleClient(GOOGLE_FOLDER)
    result = import_rows([row("https://drive.google.com/file/d/folder/view")], folder, tmp_path)
    assert result[0]["reason"] == "folder_link"
    disabled = FakeGoogleClient(GOOGLE_DOC, downloadable=False)
    result = import_rows([row("https://docs.google.com/document/d/doc/edit", 5)],
                         disabled, tmp_path)
    assert result[0]["reason"] == "download_disabled"


def test_loader_preserves_page_and_sheet_provenance(tmp_path: Path) -> None:
    content = b"%PDF-1.7\nfixture"
    (tmp_path / "abc.pdf").write_bytes(content)
    record = row("https://drive.google.com/file/d/abc/view").to_dict() | {
        "status": "downloaded", "local_path": "abc.pdf",
        "sha256": hashlib.sha256(content).hexdigest(),
        "retrieved_at": "2026-09-29T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n" + json.dumps(record) + "\n")
    docs = load_syllabus_documents(
        manifest,
        loader_factory=lambda _: [Document(page_content="Syllabus topics: transformers",
                                           metadata={"page": 2})],
    )
    assert len(docs) == 1
    assert docs[0].metadata["page"] == 3
    assert docs[0].metadata["course_code"] == "CSCI 361"
    assert docs[0].metadata["historical_source"] is True
    assert docs[0].metadata["source_url"] == record["syllabus_url"]
    expected_id = hashlib.sha256(
        f"{record['source_url']}|{record['sha256']}|3|0|pymupdf-page-v1".encode()
    ).hexdigest()
    assert docs[0].metadata["document_id"] == expected_id


def test_docx_loader_preserves_heading_table_and_citation_location(tmp_path: Path) -> None:
    content = docx_bytes()
    (tmp_path / "word.docx").write_bytes(content)
    record = row("https://drive.google.com/file/d/word/view").to_dict() | {
        "status": "downloaded", "local_path": "word.docx",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "sha256": hashlib.sha256(content).hexdigest(),
        "retrieved_at": "2026-09-30T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n")

    docs = load_syllabus_documents(manifest)

    assert len(docs) == 1
    assert "Course overview" in docs[0].page_content
    assert "Topics include computer vision." in docs[0].page_content
    assert "Assessment | Weight" in docs[0].page_content
    assert docs[0].metadata["section"] == "Course overview"
    assert "page" not in docs[0].metadata
    assert docs[0].metadata["source_url"] == record["syllabus_url"]


def test_docx_merged_cells_are_emitted_once(tmp_path: Path) -> None:
    word = WordDocument()
    table = word.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Merged heading"
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(1, 0).text = "Prerequisite"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).text = "CSCI 101"
    table.cell(1, 2).text = "CSCI 102"
    table.cell(2, 1).text = "CSCI 201"
    table.cell(2, 2).text = "CSCI 202"
    path = tmp_path / "merged.docx"
    word.save(path)
    content = path.read_bytes()
    record = row("https://drive.google.com/file/d/merged/view").to_dict() | {
        "status": "downloaded", "local_path": path.name,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "sha256": hashlib.sha256(content).hexdigest(),
        "retrieved_at": "2026-09-30T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n")

    chunks = load_syllabus_documents(manifest)
    text = "\n".join(chunk.page_content for chunk in chunks)

    assert text.count("Merged heading") == 1
    assert text.count("Prerequisite") == 1
    assert all(text.count(code) == 1 for code in ("CSCI 101", "CSCI 102", "CSCI 201", "CSCI 202"))


def test_docx_chunk_tracks_paragraph_and_table_row_positions(tmp_path: Path) -> None:
    word = WordDocument()
    word.add_paragraph("Opening paragraph")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Hours"
    table.cell(1, 0).text = "Vision"
    table.cell(1, 1).text = "10"
    word.add_paragraph("Closing paragraph")
    path = tmp_path / "blocks.docx"
    word.save(path)
    record = row("https://drive.google.com/file/d/blocks/view").to_dict() | {
        "status": "downloaded", "local_path": path.name,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "retrieved_at": "2026-09-30T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n")

    chunks = load_syllabus_documents(manifest)

    assert len(chunks) == 1
    assert chunks[0].metadata["block_start"] == 1
    assert chunks[0].metadata["block_end"] == 4
    assert chunks[0].page_content.index("Opening paragraph") < chunks[0].page_content.index("Topic | Hours")
    assert chunks[0].page_content.index("Vision | 10") < chunks[0].page_content.index("Closing paragraph")


def test_docx_long_document_chunks_on_block_boundaries(tmp_path: Path) -> None:
    word = WordDocument()
    for number in range(1, 6):
        word.add_paragraph(f"Block {number}: " + "x" * 850)
    path = tmp_path / "long.docx"
    word.save(path)
    record = row("https://drive.google.com/file/d/long/view").to_dict() | {
        "status": "downloaded", "local_path": path.name,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "retrieved_at": "2026-09-30T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n")

    chunks = load_syllabus_documents(manifest)

    assert [(chunk.metadata["block_start"], chunk.metadata["block_end"])
            for chunk in chunks] == [(1, 2), (3, 4), (5, 5)]
    assert len({chunk.metadata["document_id"] for chunk in chunks}) == 3
    assert all(len(chunk.page_content) <= 2400 for chunk in chunks)
    assert [chunk.page_content.count(f"Block {number}:") for number in range(1, 6)
            for chunk in chunks if f"Block {number}:" in chunk.page_content] == [1] * 5


def test_loader_rejects_changed_or_unsafe_snapshots(tmp_path: Path) -> None:
    record = row("https://drive.google.com/file/d/abc/view").to_dict() | {
        "status": "downloaded", "local_path": "../outside.pdf",
        "sha256": "bad", "retrieved_at": "2026-09-29T00:00:00+00:00",
    }
    manifest = tmp_path / "manifest.jsonl"
    manifest.write_text(json.dumps(record) + "\n")
    with pytest.raises(ValueError, match="Unsafe manifest path"):
        load_syllabus_documents(manifest, loader_factory=lambda _: [])
